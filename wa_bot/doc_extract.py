"""
Extractor de texto desde múltiples formatos de archivos.

Soporta: .txt .pdf .docx .xlsx .pptx

Cada extractor devuelve el texto plano. La idea es que después se procese con
parser.parse_faq_document() para extraer pares Q→A o entradas de KB.
"""
from __future__ import annotations
import logging
import io

logger = logging.getLogger(__name__)


def extract_text_from_bytes(filename: str, content: bytes) -> str | None:
    """
    Detecta formato por extensión y extrae texto plano.
    Retorna None si el formato no está soportado o falla.
    """
    if not filename or not content:
        return None
    ext = filename.lower().rsplit('.', 1)[-1] if '.' in filename else ''

    try:
        if ext == 'txt':
            try:
                return content.decode('utf-8')
            except UnicodeDecodeError:
                return content.decode('latin-1', errors='replace')

        if ext == 'pdf':
            return _extract_pdf(content)

        if ext in ('docx', 'doc'):
            return _extract_docx(content)

        if ext in ('xlsx', 'xls'):
            return _extract_xlsx(content)

        if ext == 'pptx':
            return _extract_pptx(content)

    except Exception as e:
        logger.error(f'[DocExtract] {filename} ext={ext}: {e}', exc_info=True)
        return None

    logger.warning(f'[DocExtract] formato no soportado: {filename}')
    return None


def _extract_pdf(content: bytes) -> str:
    """Extrae texto de un PDF usando pypdf."""
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(content))
    pages = []
    for i, page in enumerate(reader.pages):
        try:
            txt = page.extract_text() or ''
            if txt.strip():
                pages.append(f'--- Página {i+1} ---\n{txt}')
        except Exception as e:
            logger.debug(f'[DocExtract] PDF page {i+1}: {e}')
    return '\n\n'.join(pages)


def _extract_docx(content: bytes) -> str:
    """Extrae texto de un DOCX usando python-docx."""
    from docx import Document
    doc = Document(io.BytesIO(content))
    parts = []
    # Párrafos
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            # Detecta si es heading para preservar estructura
            style = (p.style.name or '').lower() if p.style else ''
            if 'heading' in style:
                parts.append(f'## {text}')
            else:
                parts.append(text)
    # Tablas
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(' | '.join(cells))
    return '\n\n'.join(parts)


def _extract_xlsx(content: bytes) -> str:
    """Extrae texto de un XLSX usando openpyxl."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    parts = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        parts.append(f'## Hoja: {sheet_name}')
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                parts.append(' | '.join(cells))
        parts.append('')
    return '\n'.join(parts)


def _extract_pptx(content: bytes) -> str:
    """Extrae texto de un PPTX usando python-pptx."""
    from pptx import Presentation
    prs = Presentation(io.BytesIO(content))
    parts = []
    for i, slide in enumerate(prs.slides):
        slide_parts = [f'## Slide {i+1}']
        for shape in slide.shapes:
            if hasattr(shape, 'text') and shape.text:
                slide_parts.append(shape.text.strip())
        if len(slide_parts) > 1:
            parts.append('\n'.join(slide_parts))
    return '\n\n'.join(parts)


def supported_extensions() -> list[str]:
    return ['txt', 'pdf', 'docx', 'xlsx', 'pptx']
