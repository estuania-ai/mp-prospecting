"""Rutas Flask: Campanas"""
from flask import Blueprint, request, jsonify, send_file
from database import get_db
import io

campaigns_bp = Blueprint('campaigns', __name__)


@campaigns_bp.route('/', methods=['GET'])
def get_campaigns():
    conn = get_db()
    rows = conn.execute('''
        SELECT * FROM campaigns ORDER BY sent_at DESC LIMIT 50
    ''').fetchall()
    conn.close()
    return jsonify([dict(r) for r in rows])


@campaigns_bp.route('/messages', methods=['GET'])
def get_messages():
    conn = get_db()
    limit = int(request.args.get('limit', 100))
    tipo  = request.args.get('type', '')
    q = 'SELECT * FROM messages WHERE 1=1'
    params = []
    if tipo:
        q += ' AND message_type = ?'
        params.append(tipo)
    q += ' ORDER BY sent_at DESC LIMIT ?'
    params.append(limit)
    rows = conn.execute(q, params).fetchall()
    conn.close()
    return jsonify([dict(r) for r in rows])


@campaigns_bp.route('/<int:campaign_id>/export', methods=['GET'])
def export_campaign(campaign_id):
    """Exporta los mensajes de una campana como Excel"""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        return jsonify({'error': 'openpyxl no instalado. Ejecuta: pip install openpyxl'}), 500

    conn = get_db()

    # Info de la campana
    campaign = conn.execute('SELECT * FROM campaigns WHERE id = ?', (campaign_id,)).fetchone()

    # Si campaign_id = 0 o no existe, exportar todos los mensajes
    if campaign:
        # Buscar mensajes de esa campana por fecha aproximada
        sent_date = dict(campaign)['sent_at'][:10] if campaign else None
        if sent_date:
            messages = conn.execute('''
                SELECT
                    m.id,
                    l.name as negocio,
                    m.phone,
                    l.comuna,
                    l.rubro,
                    m.message_type as tipo,
                    m.sent_at as fecha_envio,
                    CASE m.status
                        WHEN 'sent' THEN 'Enviado'
                        WHEN 'failed' THEN 'Fallido'
                        WHEN 'phone_not_exists' THEN 'Numero no existe'
                        ELSE m.status
                    END as estado_envio,
                    CASE ls.status
                        WHEN 'enviado' THEN 'Enviado'
                        WHEN 'abierto' THEN 'Abierto'
                        WHEN 'interesado' THEN 'Interesado'
                        WHEN 'quiere_reunion' THEN 'Quiere reunion'
                        WHEN 'cerrado' THEN 'Cerrado'
                        WHEN 'no_interesado' THEN 'No interesado'
                        WHEN 'opt_out' THEN 'Opt-out'
                        WHEN 'telefono_no_existe' THEN 'Numero no existe'
                        ELSE 'Sin estado'
                    END as estado_lead,
                    m.opened_at as fecha_apertura
                FROM messages m
                LEFT JOIN leads l ON m.lead_id = l.id
                LEFT JOIN lead_status ls ON l.id = ls.lead_id
                WHERE date(m.sent_at) = ?
                ORDER BY m.sent_at ASC
            ''', (sent_date,)).fetchall()
        else:
            messages = []
    else:
        messages = conn.execute('''
            SELECT
                m.id,
                l.name as negocio,
                m.phone,
                l.comuna,
                l.rubro,
                m.message_type as tipo,
                m.sent_at as fecha_envio,
                CASE m.status
                    WHEN 'sent' THEN 'Enviado'
                    WHEN 'failed' THEN 'Fallido'
                    WHEN 'phone_not_exists' THEN 'Numero no existe'
                    ELSE m.status
                END as estado_envio,
                CASE ls.status
                    WHEN 'enviado' THEN 'Enviado'
                    WHEN 'abierto' THEN 'Abierto'
                    WHEN 'interesado' THEN 'Interesado'
                    WHEN 'quiere_reunion' THEN 'Quiere reunion'
                    WHEN 'cerrado' THEN 'Cerrado'
                    WHEN 'no_interesado' THEN 'No interesado'
                    WHEN 'opt_out' THEN 'Opt-out'
                    WHEN 'telefono_no_existe' THEN 'Numero no existe'
                    ELSE 'Sin estado'
                END as estado_lead,
                m.opened_at as fecha_apertura
            FROM messages m
            LEFT JOIN leads l ON m.lead_id = l.id
            LEFT JOIN lead_status ls ON l.id = ls.lead_id
            ORDER BY m.sent_at DESC
            LIMIT 500
        ''').fetchall()

    conn.close()

    # ── Crear Excel ──────────────────────────────────────
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Prospeccion"

    # Colores corporativos MP
    COLOR_BLUE  = "009EE3"
    COLOR_YELLOW= "FFE600"
    COLOR_WHITE = "FFFFFF"
    COLOR_GRAY  = "F0F4F8"
    COLOR_GREEN = "00A650"
    COLOR_RED   = "F23D4F"
    COLOR_WARN  = "FF9500"

    # Header titulo
    ws.merge_cells('A1:J1')
    title_cell = ws['A1']
    title_cell.value = f"MercadoPago POS - Reporte de Prospeccion"
    title_cell.font = Font(bold=True, size=14, color=COLOR_WHITE)
    title_cell.fill = PatternFill("solid", fgColor=COLOR_BLUE)
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30

    # Subtitulo con nombre campana
    ws.merge_cells('A2:J2')
    sub_cell = ws['A2']
    camp_name = dict(campaign)['name'] if campaign else 'Todos los mensajes'
    sub_cell.value = camp_name
    sub_cell.font = Font(bold=True, size=11, color="333333")
    sub_cell.fill = PatternFill("solid", fgColor=COLOR_YELLOW)
    sub_cell.alignment = Alignment(horizontal='center')
    ws.row_dimensions[2].height = 22

    # Encabezados columnas
    headers = [
        ('N°', 5),
        ('Negocio', 30),
        ('Telefono', 16),
        ('Comuna', 18),
        ('Rubro', 18),
        ('Tipo Envio', 14),
        ('Fecha Envio', 20),
        ('Estado Envio', 16),
        ('Estado Lead', 16),
        ('Fecha Apertura', 20),
    ]

    header_fill   = PatternFill("solid", fgColor=COLOR_BLUE)
    header_font   = Font(bold=True, color=COLOR_WHITE, size=10)
    header_align  = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin_border   = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )

    for col, (header, width) in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=header)
        cell.font   = header_font
        cell.fill   = header_fill
        cell.alignment = header_align
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[3].height = 28

    # Filas de datos
    fill_alt  = PatternFill("solid", fgColor=COLOR_GRAY)
    fill_white= PatternFill("solid", fgColor=COLOR_WHITE)
    font_data = Font(size=10)
    align_center = Alignment(horizontal='center', vertical='center')
    align_left   = Alignment(horizontal='left',   vertical='center')

    # Colores por estado
    estado_colors = {
        'Enviado':        ('E8F4FD', '0078C8'),
        'Abierto':        ('E8F3FF', '3D5A99'),
        'Interesado':     ('FFF3CD', '856404'),
        'Quiere reunion': ('EDE9FE', '5B21B6'),
        'Cerrado':        ('D4EDDA', '155724'),
        'No interesado':  ('FDE8EA', 'A01622'),
        'Opt-out':        ('F0F0F0', '666666'),
        'Numero no existe':('FFF3E0','E65100'),
        'Fallido':        ('FDE8EA', 'A01622'),
        'Sin estado':     ('F8F9FA', '6C757D'),
    }

    for i, row in enumerate(messages):
        row_dict = dict(row)
        excel_row = i + 4
        is_alt = i % 2 == 0
        row_fill = fill_alt if is_alt else fill_white

        estado_e = row_dict.get('estado_envio', 'Sin estado')
        estado_l = row_dict.get('estado_lead',  'Sin estado')

        values = [
            i + 1,
            row_dict.get('negocio', ''),
            row_dict.get('phone', ''),
            row_dict.get('comuna', ''),
            row_dict.get('rubro', ''),
            row_dict.get('tipo', ''),
            row_dict.get('fecha_envio', ''),
            estado_e,
            estado_l,
            row_dict.get('fecha_apertura', ''),
        ]

        for col, value in enumerate(values, 1):
            cell = ws.cell(row=excel_row, column=col, value=value)
            cell.font   = font_data
            cell.border = thin_border

            # Color segun estado
            if col == 8 and estado_e in estado_colors:
                bg, fg = estado_colors[estado_e]
                cell.fill = PatternFill("solid", fgColor=bg)
                cell.font = Font(size=10, color=fg, bold=True)
                cell.alignment = align_center
            elif col == 9 and estado_l in estado_colors:
                bg, fg = estado_colors[estado_l]
                cell.fill = PatternFill("solid", fgColor=bg)
                cell.font = Font(size=10, color=fg, bold=True)
                cell.alignment = align_center
            elif col in (1, 3, 6):
                cell.fill = row_fill
                cell.alignment = align_center
            else:
                cell.fill = row_fill
                cell.alignment = align_left

        ws.row_dimensions[excel_row].height = 18

    # Fila resumen final
    total_row = len(messages) + 4
    ws.merge_cells(f'A{total_row}:G{total_row}')
    summary_cell = ws.cell(row=total_row, column=1, value=f'Total: {len(messages)} mensajes')
    summary_cell.font  = Font(bold=True, size=10, color=COLOR_WHITE)
    summary_cell.fill  = PatternFill("solid", fgColor=COLOR_BLUE)
    summary_cell.alignment = align_center

    # Contar por estado
    estados_count = {}
    for row in messages:
        e = dict(row).get('estado_lead','Sin estado')
        estados_count[e] = estados_count.get(e, 0) + 1

    summary_text = ' | '.join([f"{k}: {v}" for k,v in estados_count.items()])
    ws.merge_cells(f'H{total_row}:J{total_row}')
    sum2 = ws.cell(row=total_row, column=8, value=summary_text)
    sum2.font  = Font(bold=True, size=9, color=COLOR_WHITE)
    sum2.fill  = PatternFill("solid", fgColor=COLOR_BLUE)
    sum2.alignment = align_center
    ws.row_dimensions[total_row].height = 20

    # Guardar en memoria
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Campana_{campaign_id}_{camp_name[:30].replace(' ','_')}.xlsx"

    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename
    )
